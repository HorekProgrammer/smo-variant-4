"""Build local PlantUML images and review documents. No publication operations."""

from __future__ import annotations

import argparse
from html import escape
import math
from pathlib import Path
import re
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DIAGRAMS = DOCS / "diagrams"
IMAGES = DOCS / "img"
TITLES = {
    "sequence": "Общий сиквенс",
    "sequence-happy": "HappyFlow: свободные буфер и прибор",
    "sequence-overflow": "Переполнение, вытеснение и последующая обработка",
    "sequence-selection": "Выбор заявки по приоритету и прибора по кольцу",
    "classes": "Диаграмма классов",
    "flowchart": "Блок-схема с дорожками",
    "archimate": "Бизнес-процесс ArchiMate 3.0",
}


def render(jar: Path, java: str) -> None:
    sources = [str(DIAGRAMS / f"{name}.puml") for name in TITLES]
    base = [java, "-Djava.awt.headless=true", "-Xmx1024m", "-jar", str(jar),
            "-charset", "UTF-8", "-failfast2"]
    subprocess.run(base + ["-checkonly"] + sources, check=True, cwd=ROOT)
    IMAGES.mkdir(exist_ok=True)
    for kind in ("svg", "png"):
        subprocess.run(base + [f"-t{kind}", "-o", str(IMAGES)] + sources,
                       check=True, cwd=ROOT)
        for name in TITLES:
            target = IMAGES / f"{name}.{kind}"
            if not target.is_file() or target.stat().st_size == 0:
                raise RuntimeError(f"Missing rendered diagram: {target}")
            if kind == "svg" and "Syntax Error" in target.read_text(encoding="utf-8"):
                raise RuntimeError(f"PlantUML error in {target}")


def build_gallery() -> None:
    cards = []
    for name, title in TITLES.items():
        cards.append(f'''<section id="{name}"><h2>{escape(title)}</h2>
<p><a href="diagrams/{name}.puml">PlantUML</a> ·
<a href="img/{name}.svg" target="_blank">Открыть SVG и увеличить</a> ·
<a href="img/{name}.png" target="_blank">PNG</a></p>
<a href="img/{name}.svg" target="_blank"><img src="img/{name}.svg"
alt="{escape(title)}" loading="lazy"></a></section>''')
    nav = " · ".join(f'<a href="#{name}">{escape(title)}</a>'
                     for name, title in TITLES.items())
    page = '''<!doctype html><html lang="ru"><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>СМО — вариант 4 · Первичные артефакты</title>
<style>
body{margin:0;background:#eef2f6;color:#18283a;font:17px/1.6 Arial,sans-serif}
main{max-width:1400px;margin:auto;padding:32px}
header,section{background:white;padding:28px;margin:0 0 24px;border-radius:12px;
box-shadow:0 2px 12px #17283c0b}h1{line-height:1.2}h2{margin-top:0}
a{color:#115c9f}nav{font-size:14px}img{max-width:100%;height:auto;border:1px solid #dde5ed}
.code{padding:12px;background:#edf5fc;border-radius:6px;font-weight:bold}
@media print{body{background:white}main{padding:0}section{break-before:page;box-shadow:none}}
</style><main><header><h1>Сервис обработки аналитических заданий спортивных клубов</h1>
<p>Первичные артефакты · Вариант 4<br>Добров М. Г. · группа 5130904/40106</p>
<p class="code">ИБ · ИЗ2 · ПЗ1 · Д1ОЗ1 · Д1ОО4 · Д2П2 · Д2Б4 · ОР2 · ОД1</p>
<p>Три клуба → приём → общий кольцевой буфер → выбор по приоритету клуба →
свободный обработчик по кольцу → аналитический отчёт.</p>
<p>Все заявки проходят через буфер. При переполнении вытесняется последняя старая
ожидающая заявка. Среди заявок одного клуба выбирается последняя поступившая.</p>
<p><a href="primary-artifacts.md">Подробное описание Markdown</a> ·
<a href="review-checklist.md">Контрольные сценарии</a></p>
<p>Нажмите на диаграмму, чтобы открыть масштабируемый SVG. Галерея работает локально.</p>
<nav>''' + nav + "</nav></header>" + "\n".join(cards) + "</main></html>"
    (DOCS / "index.html").write_text(page, encoding="utf-8")


def plain(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


def add_markdown(doc: Document, source: Path, page_start: bool = False) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if match:
                doc.add_paragraph(f"Рисунок: {match[1]}. См. приложение с диаграммами и HTML-галерею.")
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [plain(c.strip()) for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                table = doc.add_table(rows=0, cols=len(rows[0]))
                table.style = "Light Shading Accent 1"
                for cells in rows:
                    row = table.add_row().cells
                    for cell, text in zip(row, cells):
                        cell.text = text
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(3)
                            paragraph.paragraph_format.line_spacing = 1
                            for run in paragraph.runs:
                                run.font.size = Pt(9)
                for cell in table.rows[0].cells:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
            continue
        if line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 3)
            heading = doc.add_heading(plain(line.lstrip("# ")), level=level)
            if page_start:
                heading.paragraph_format.page_break_before = True
                page_start = False
        elif line.startswith("- "):
            doc.add_paragraph(plain(line[2:]), style="List Bullet")
        else:
            doc.add_paragraph(plain(line))
        i += 1


def build_word() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    doc.core_properties.title = "Первичные артефакты СМО. Вариант 4"
    doc.core_properties.author = "Добров М. Г."
    doc.add_heading("ПЕРВИЧНЫЕ АРТЕФАКТЫ", 0)
    doc.add_heading("Сервис обработки аналитических заданий спортивных клубов", 1)
    doc.add_paragraph("Вариант 4\nДобров М. Г.\nГруппа 5130904/40106")
    doc.add_paragraph("ИБ · ИЗ2 · ПЗ1 · Д1ОЗ1 · Д1ОО4 · Д2П2 · Д2Б4 · ОР2 · ОД1")
    doc.add_paragraph("Этап проектирования. Описание, диаграммы и контрольные сценарии.")
    doc.add_paragraph("Полноразмерные масштабируемые диаграммы: docs/index.html и docs/img/*.svg. "
                      "В приложении длинные схемы разделены на листы с перекрытием.")
    add_markdown(doc, DOCS / "primary-artifacts.md", page_start=True)
    add_markdown(doc, DOCS / "mapping.md", page_start=True)
    add_markdown(doc, DOCS / "review-checklist.md", page_start=True)

    # Landscape A3 appendix keeps wide diagrams readable, while tall diagrams
    # continue vertically across pages. Original PNG and SVG stay intact.
    appendix = doc.add_section(WD_SECTION_START.NEW_PAGE)
    appendix.page_width, appendix.page_height = Cm(42), Cm(29.7)
    appendix.top_margin = appendix.bottom_margin = Cm(1.4)
    appendix.left_margin = appendix.right_margin = Cm(1.5)
    crop_dir = IMAGES / "word-parts"
    crop_dir.mkdir(exist_ok=True)
    first = True
    for name, title in TITLES.items():
        with Image.open(IMAGES / f"{name}.png") as original:
            # A smaller image can retain its natural size; wide images fill 39cm.
            width_cm = min(39, original.width / 96 * 2.54)
            pixels_per_cm = original.width / width_cm
            part_height = max(100, int(23.5 * pixels_per_cm))
            overlap = min(60, part_height // 12)
            count = max(1, math.ceil((original.height - overlap) / (part_height - overlap)))
            part_height = math.ceil((original.height + (count - 1) * overlap) / count)
            top, part = 0, 1
            while top < original.height:
                if not first:
                    doc.add_page_break()
                first = False
                bottom = min(original.height, top + part_height)
                suffix = f" · лист {part}" if original.height > part_height else ""
                doc.add_heading(f"Приложение. {title}{suffix}", 1)
                doc.add_paragraph(f"Исходник: docs/diagrams/{name}.puml · Полная схема: docs/img/{name}.svg")
                crop = original.crop((0, top, original.width, bottom))
                target = crop_dir / f"{name}-{part}.png"
                crop.save(target)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(target), width=Cm(width_cm))
                if bottom == original.height:
                    break
                top = bottom - overlap
                part += 1
    doc.save(DOCS / "primary-artifacts.docx")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--plantuml-jar", type=Path, required=True)
    parser.add_argument("--java", default="java", help="Path to Java 11+ executable")
    args = parser.parse_args()
    jar = args.plantuml_jar.resolve()
    if not jar.is_file():
        parser.error(f"PlantUML jar not found: {jar}")
    render(jar, args.java)
    build_gallery()
    build_word()
    print(f"Built {len(TITLES)} diagrams (SVG + PNG), HTML gallery and Word document in {DOCS}")


if __name__ == "__main__":
    main()
